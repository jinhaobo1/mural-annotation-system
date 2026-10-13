# -*- coding: utf-8 -*-
"""
壁画展示系统 · 标注结果导出工具（多格式版 · 便携修复版）
功能：从数据库读取标注结果，导出为 JSON / CSV / Excel 三种格式
用法：
    双击运行（默认导出全部三种格式）
    命令行指定格式： python exporter.py json
                     python exporter.py csv
                     python exporter.py excel
                     python exporter.py all        （默认）
    指定输出目录：   python exporter.py all D:\\output
路径查找规则（便携式，不再写死绝对路径）：
    1. PyInstaller 解包目录（sys._MEIPASS，数据库打包进 exe 时生效）
    2. 程序所在目录（exe 或 .py 文件所在文件夹）
    3. 当前工作目录
"""
import os
import sys
import csv
import json
import sqlite3
from datetime import datetime
import io
# 强制标准输出为 UTF-8，避免 GBK 控制台报错
if sys.stdout and sys.stdout.encoding and sys.stdout.encoding.lower() not in ("utf-8", "utf8"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")
# ==================== 便携式路径解析 ====================
def app_dir():
    """程序所在目录（exe 与 .py 通用）"""
    if getattr(sys, "frozen", False):
        # PyInstaller 打包后：sys.executable 是 exe 的完整路径
        return os.path.dirname(os.path.abspath(sys.executable))
    # 普通 .py 运行：__file__ 是脚本路径
    return os.path.dirname(os.path.abspath(__file__))
def bundled_dir():
    """PyInstaller 解包资源目录（--onefile 运行时存在；未打包时为 None）"""
    return getattr(sys, "_MEIPASS", None)
# ==================== 配置区 ====================
DB_NAME = "壁画数据库.db"
# 数据库查找优先级：
#   1. exe 内嵌资源（--add-data 打包进去的副本）
#   2. 程序所在文件夹（exe/py 旁边）
#   3. 当前工作目录
DB_CANDIDATES = []
_bd = bundled_dir()
if _bd:
    DB_CANDIDATES.append(os.path.join(_bd, DB_NAME))
DB_CANDIDATES.append(os.path.join(app_dir(), DB_NAME))
DB_CANDIDATES.append(os.path.join(os.getcwd(), DB_NAME))
# 默认输出目录：程序所在文件夹（不再写死桌面）
DEFAULT_OUTPUT_DIR = app_dir()
# =============================================
def find_db():
    """自动查找数据库文件，返回第一个存在的路径"""
    seen = set()
    for p in DB_CANDIDATES:
        if p in seen:
            continue
        seen.add(p)
        if os.path.isfile(p):
            return p
    return None
def read_annotations(db_path):
    """从数据库读取全部标注结果"""
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            "SELECT id, name, dynasty, description, site, category, "
            "confidence, image_name, image_path, image_exists, created_at "
            "FROM murals ORDER BY id"
        ).fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()
def to_structured(rows):
    """整理为结构化数据（统一内部结构）"""
    meta = {
        "title": "壁画标注结果",
        "source": "壁画展示系统 · SQLite 数据库",
        "exported_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "total": len(rows),
    }
    data = []
    for r in rows:
        data.append({
            "id": r["id"],
            "name": r["name"],
            "dynasty": r["dynasty"],
            "description": r["description"],
            "site": r["site"],
            "category": r["category"],
            "confidence": r["confidence"],
            "image_name": r["image_name"],
            "image_path": r["image_path"],
            "image_exists": bool(r["image_exists"]),
        })
    return meta, data
def export_json(meta, data, out_dir):
    """导出 JSON"""
    out = os.path.join(out_dir, "壁画标注结果.json")
    payload = {
        "meta": meta,
        "data": [
            {
                "id": d["id"],
                "name": d["name"],
                "dynasty": d["dynasty"],
                "description": d["description"],
                "site": d["site"],
                "category": d["category"],
                "confidence": d["confidence"],
                "image": {
                    "file_name": d["image_name"],
                    "path": d["image_path"],
                    "exists": d["image_exists"],
                },
            }
            for d in data
        ],
    }
    with open(out, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    return out
def export_csv(meta, data, out_dir):
    """导出 CSV（UTF-8 BOM，Excel 直接打开不乱码）"""
    out = os.path.join(out_dir, "壁画标注结果.csv")
    headers = ["id", "name", "dynasty", "description", "site",
               "category", "confidence", "image_name", "image_path", "image_exists"]
    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for d in data:
            w.writerow([d[h] if h != "image_exists" else ("是" if d[h] else "否")
                        for h in headers])
    return out
def export_excel(meta, data, out_dir):
    """导出 Excel（带表头样式与列宽）"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise ImportError("缺少 openpyxl 库，请先安装：pip install openpyxl")
    out = os.path.join(out_dir, "壁画标注结果.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "壁画标注结果"
    headers = ["ID", "名称", "朝代", "简介", "所属石窟",
               "分类", "把握度", "图片文件名", "图片路径", "图片存在"]
    keys = ["id", "name", "dynasty", "description", "site",
            "category", "confidence", "image_name", "image_path", "image_exists"]
    head_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="764BA2")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    thin = Side(style="thin", color="D0C4E0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = center
        cell.border = border
    for r, d in enumerate(data, 2):
        for c, k in enumerate(keys, 1):
            val = d[k]
            if k == "image_exists":
                val = "是" if val else "否"
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = border
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = left if k == "description" else center
        if r % 2 == 0:
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="F3EEF9")
    widths = [6, 18, 12, 60, 18, 12, 8, 16, 42, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 26
    wb.save(out)
    return out
def main():
    print("=" * 60)
    print("   壁画展示系统 · 标注结果导出工具（多格式版）")
    print("=" * 60)
    # 解析参数
    fmt = "all"
    out_dir = DEFAULT_OUTPUT_DIR
    args = sys.argv[1:]
    if args:
        fmt = args[0].lower()
    if len(args) > 1:
        out_dir = args[1]
    # 1. 查找数据库（便携式多路径）
    db = find_db()
    if not db:
        print("[错误] 未找到数据库文件「%s」，已尝试以下位置：" % DB_NAME)
        for p in DB_CANDIDATES:
            print("   -", p)
        print("请将数据库文件与程序放在同一文件夹后重试。")
        input("\n按回车键退出...")
        return
    print("数据库：", db)
    if getattr(sys, "_MEIPASS", None) and db.startswith(sys._MEIPASS):
        print("（使用 exe 内嵌的数据库副本）")
    # 2. 读取
    try:
        rows = read_annotations(db)
    except Exception as e:
        print("[错误] 读取失败：", e)
        input("\n按回车键退出...")
        return
    if not rows:
        print("[警告] 数据库中没有标注数据")
        input("\n按回车键退出...")
        return
    print("读取到", len(rows), "条标注结果")
    meta, data = to_structured(rows)
    os.makedirs(out_dir, exist_ok=True)
    print("输出目录：", out_dir)
    # 3. 按格式导出
    results = []
    tasks = []
    if fmt in ("all", "json"):
        tasks.append(("JSON", export_json))
    if fmt in ("all", "csv"):
        tasks.append(("CSV", export_csv))
    if fmt in ("all", "excel", "xlsx", "xls"):
        tasks.append(("Excel", export_excel))
    if not tasks:
        print("[警告] 未知格式：", fmt, "（可选 json / csv / excel / all）")
        input("\n按回车键退出...")
        return
    for label, func in tasks:
        try:
            path = func(meta, data, out_dir)
            size = round(os.path.getsize(path) / 1024, 1)
            print("[成功]", label, "已导出：", path, "（%sKB）" % size)
            results.append((label, path, True))
        except Exception as e:
            print("[失败]", label, "导出失败：", e)
            results.append((label, None, False))
    # 4. 预览
    print("\n" + "-" * 60)
    print("%-4s%-18s%-12s%s" % ("ID", "名称", "朝代", "图片"))
    print("-" * 60)
    for d in data:
        flag = "Y" if d["image_exists"] else "X"
        print("%-4d%-18s%-12s%s" % (d["id"], d["name"], d["dynasty"], flag))
    print("-" * 60)
    ok = sum(1 for _, _, s in results if s)
    print("共", meta["total"], "条 | 成功导出", ok, "/", len(results), "种格式 |", meta["exported_at"])
    print("导出完成！")
    input("\n按回车键退出...")
if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        import traceback
        traceback.print_exc()
        input("\n程序出错，按回车键退出...")